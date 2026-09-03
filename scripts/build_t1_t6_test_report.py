from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table


ROOT = Path(__file__).resolve().parents[1]
OUTLINE = Path(r"C:\Users\Dbook\Desktop\中期材料\测试大纲20260818V4.docx")
TEMPLATE = Path(r"C:\Users\Dbook\Desktop\中期材料\测试报告_lzy.docx")
OUTPUT = ROOT / "docs" / "测试报告_T1-T6_20260819_V4.docx"
ASSET_DIR = ROOT / "tmp_docx" / "test_report_assets"

CHINESE_FONT = "仿宋_GB2312"
LATIN_FONT = "Times New Roman"
CHART_FONT_PATH = r"C:\Windows\Fonts\simfang.ttf"
font_manager.fontManager.addfont(CHART_FONT_PATH)
plt.rcParams["font.family"] = font_manager.FontProperties(
    fname=CHART_FONT_PATH).get_name()
plt.rcParams["axes.unicode_minus"] = False


CASES = [
    {
        "id": "T-1",
        "model": "",
        "dataset": "",
        "title": "T-1 深度学习模型跨域协同平台并发性能测试",
        "summary": "建立并保持10000个连接，最低QPS为71677.59；新增客户端模型参数往返成功，服务端错误数为0。",
        "actual": (
            "实际建立并持续保持10000个客户端连接，观测时10000个模型参数事务均处于活动状态，"
            "10000次模型传输全部完成；最低实测QPS为71677.59，高于10000；新增客户端成功完成模型参数获取与回传，"
            "往返耗时0.302秒，服务端错误数为0。用例总耗时406.964秒，测试通过。\n"
            "结果文件：docs/test_logs/concurrent_availability/outline_t01/aggregate_summary.json。"
        ),
    },
    {
        "id": "T-2",
        "model": "ViT",
        "dataset": "MDDigits",
        "title": "T-2 ViT模型跨域训练支持与准确率测试",
        "logs": {
            "FedAvg": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t2_digit3_vit_formal_fedavg.log",
            "FedProx": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t2_digit3_vit_formal_fedprox.log",
            "平台": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t2_digit3_vit_formal_platform.log",
        },
        "accuracies": {"FedAvg": 0.3993, "FedProx": 0.4473, "平台": 0.7927},
        "gains": (39.34, 34.54),
        "times": (56.690, 67.967, 175.597),
        "config": "训练轮数100、客户端数60、批大小32、每轮本地更新5步、Adam优化器、学习率0.0001、狄利克雷参数集中度α=0.1",
        "task_file": "scripts/test_outline_validation/task_profiles/mddigits_class_targets.json",
        "distribution_dir": "exp/digit_three_domain/training_distributions",
        "distribution": "生成60份训练数据分布文件，类别0至9的class_counts均为40，与任务文件要求一致。",
        "result_dir": "docs/test_logs/t2t6_formal_final_20260819",
    },
    {
        "id": "T-3",
        "model": "CNN",
        "dataset": "OfficeHome",
        "title": "T-3 CNN模型跨域训练支持与准确率测试",
        "logs": {
            "FedAvg": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t3_officehome_cnn_formal_fedavg.log",
            "FedProx": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t3_officehome_cnn_formal_fedprox.log",
            "平台": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t3_officehome_cnn_formal_platform.log",
        },
        "accuracies": {"FedAvg": 0.4513, "FedProx": 0.4712, "平台": 0.7525},
        "gains": (30.12, 28.13),
        "times": (23.977, 28.795, 395.530),
        "config": "训练轮数100、客户端数60、批大小32、每轮本地更新3步、Adam优化器、学习率0.0001、狄利克雷参数集中度α=0.1",
        "task_file": "scripts/test_outline_validation/task_profiles/officehome_accuracy_class_targets.json",
        "distribution_dir": "exp/test_outline_validation/T-03/training_distributions",
        "distribution": "生成60份训练数据分布文件，类别0至64的class_counts均为20，与任务文件要求一致。",
        "result_dir": "docs/test_logs/t2t6_formal_final_20260819",
    },
    {
        "id": "T-4",
        "model": "MLP",
        "dataset": "OfficeHome",
        "title": "T-4 MLP模型跨域训练支持与准确率测试",
        "logs": {
            "FedAvg": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t4_officehome_mlp_formal_fedavg.log",
            "FedProx": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t4_officehome_mlp_formal_fedprox.log",
            "平台": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t4_officehome_mlp_formal_platform.log",
        },
        "accuracies": {"FedAvg": 0.2776, "FedProx": 0.2884, "平台": 0.6187},
        "gains": (34.11, 33.03),
        "times": (22.963, 27.822, 347.917),
        "config": "训练轮数100、客户端数60、批大小32、每轮本地更新3步、Adam优化器、学习率0.0001、狄利克雷参数集中度α=0.1",
        "task_file": "scripts/test_outline_validation/task_profiles/officehome_accuracy_class_targets.json",
        "distribution_dir": "exp/test_outline_validation/T-04/training_distributions",
        "distribution": "生成60份训练数据分布文件，类别0至64的class_counts均为20，与任务文件要求一致。",
        "result_dir": "docs/test_logs/t2t6_formal_final_20260819",
    },
    {
        "id": "T-5",
        "model": "RNN",
        "dataset": "MDSent",
        "title": "T-5 RNN模型跨域训练支持与准确率测试",
        "logs": {
            "FedAvg": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t5_mdsent_rnn_formal_fedavg.log",
            "FedProx": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t5_mdsent_rnn_formal_fedprox.log",
            "平台": ROOT / "docs/test_logs/t2t6_formal_final_20260819/t5_mdsent_rnn_formal_platform.log",
        },
        "accuracies": {"FedAvg": 0.3683, "FedProx": 0.3498, "平台": 0.6019},
        "gains": (23.36, 25.21),
        "times": (5.158, 7.195, 4.918),
        "config": "训练轮数100、客户端数120、批大小32、每轮本地更新2步、Adam优化器、学习率0.0001、狄利克雷参数集中度α=0.01",
        "task_file": "scripts/test_outline_validation/task_profiles/mdsent_class_targets.json",
        "distribution_dir": "exp/test_outline_validation/T-05/training_distributions",
        "distribution": "生成3份本轮参与客户端的训练数据分布文件，类别0至3的class_counts均为50，与任务文件要求一致。",
        "result_dir": "docs/test_logs/t2t6_formal_final_20260819",
    },
    {
        "id": "T-6",
        "model": "LSTM",
        "dataset": "MDSent",
        "title": "T-6 LSTM模型跨域训练支持与准确率测试",
        "logs": {
            "FedAvg": ROOT / "docs/test_logs/t6_retuned_final_20260819/t6_mdsent_lstm_formal_fedavg.log",
            "FedProx": ROOT / "docs/test_logs/t6_retuned_final_20260819/t6_mdsent_lstm_formal_fedprox.log",
            "平台": ROOT / "docs/test_logs/t6_retuned_final_20260819/t6_mdsent_lstm_formal_platform.log",
        },
        "accuracies": {"FedAvg": 0.3862, "FedProx": 0.3716, "平台": 0.6647},
        "gains": (27.85, 29.31),
        "times": (6.369, 8.669, 8.954),
        "config": "训练轮数100、客户端数120、批大小32、每轮本地更新2步、Adam优化器、学习率0.0001、狄利克雷参数集中度α=0.01",
        "task_file": "scripts/test_outline_validation/task_profiles/mdsent_class_targets.json",
        "distribution_dir": "exp/test_outline_validation/T-06/training_distributions",
        "distribution": "生成6份本轮参与客户端的训练数据分布文件，类别0至3的class_counts均为50，与任务文件要求一致。",
        "result_dir": "docs/test_logs/t6_retuned_final_20260819",
    },
]


ROUND_PATTERN = re.compile(
    r"Server:\s*Round\s+(\d+)\s+MLP Test Accuracy.*?average:\s*([0-9.]+)")


def parse_rounds(path: Path) -> list[float]:
    found: dict[int, float] = {}
    with path.open("r", encoding="utf-8", errors="replace") as stream:
        for line in stream:
            match = ROUND_PATTERN.search(line)
            if match:
                found[int(match.group(1))] = float(match.group(2))
    if sorted(found) != list(range(100)):
        raise ValueError(f"{path} does not contain rounds 0..99 exactly: {sorted(found)}")
    return [found[i] for i in range(100)]


def make_chart(case: dict) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    chart = ASSET_DIR / f"{case['id'].replace('-', '').lower()}_accuracy_curve.png"
    styles = {
        "FedAvg": dict(color="#1F77B4", linestyle="--", linewidth=2.0,
                       marker="o", markevery=10, markersize=3.8),
        "FedProx": dict(color="#FF7F0E", linestyle="-.", linewidth=2.0,
                        marker="s", markevery=10, markersize=3.8),
        "平台": dict(color="#D62728", linestyle="-", linewidth=2.5,
                     marker="^", markevery=10, markersize=4.2),
    }
    fig, ax = plt.subplots(figsize=(7.2, 4.25), dpi=180)
    for label in ("FedAvg", "FedProx", "平台"):
        values = parse_rounds(case["logs"][label])
        ax.plot(range(100), [v * 100 for v in values], label=label, **styles[label])
    ax.set_xlim(0, 99)
    ax.set_xticks([0, 20, 40, 60, 80, 99])
    ax.set_xlabel("训练轮次")
    ax.set_ylabel("平均测试准确率（%）")
    ax.set_title(f"{case['dataset']}数据集{case['model']}模型")
    ax.grid(True, color="#d9d9d9", linewidth=0.7, alpha=0.8)
    ax.legend(loc="best", frameon=True, facecolor="white", framealpha=0.95)
    for spine in ax.spines.values():
        spine.set_color("#444444")
    fig.tight_layout()
    fig.savefig(chart, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return chart


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.size = Pt(10.5)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def replace_table(target, replacement_xml) -> None:
    target._element.getparent().replace(target._element, copy.deepcopy(replacement_xml))


def find_result_row(table: Table):
    for row in table.rows:
        if row.cells[0].text.strip() == "实际结果":
            return row._tr
    raise ValueError("Result row template not found")


def find_person_row(table: Table):
    for row in table.rows:
        if row.cells[0].text.strip() == "用例执行人员":
            return row._tr
    raise ValueError("Personnel row template not found")


def set_report_result_cell(cell, case: dict, chart: Path | None) -> None:
    # Retain the first paragraph's table style and remove prior content/images.
    first = cell.paragraphs[0]
    for paragraph in list(cell.paragraphs[1:]):
        cell._tc.remove(paragraph._p)
    for run in list(first.runs):
        first._p.remove(run._r)
    first.style = "表格正文"
    # Result records contain explicit line breaks. Left alignment keeps each
    # evidence line readable and avoids Word stretching short label lines.
    first.alignment = WD_ALIGN_PARAGRAPH.LEFT
    first.paragraph_format.first_line_indent = Pt(0)
    first.paragraph_format.left_indent = Pt(0)
    first.paragraph_format.right_indent = Pt(0)
    if case["id"] == "T-1":
        result_text = case["actual"]
    else:
        acc = case["accuracies"]
        gain_avg, gain_prox = case["gains"]
        t_avg, t_prox, t_platform = case["times"]
        total_time = t_avg + t_prox + t_platform
        result_text = (
            f"平台成功读取任务文件，{case['distribution']}平台、FedAvg和FedProx均完成100轮训练，"
            "日志连续记录第0至99轮平均测试准确率，训练过程未出现异常。"
            f"末轮平均测试准确率分别为FedAvg {acc['FedAvg'] * 100:.2f}%、"
            f"FedProx {acc['FedProx'] * 100:.2f}%、平台{acc['平台'] * 100:.2f}%。"
            f"平台较FedAvg和FedProx分别提高{gain_avg:.2f}%和{gain_prox:.2f}%，均不小于20%。"
            f"三种方法执行时间分别为{t_avg:.3f}秒、{t_prox:.3f}秒和{t_platform:.3f}秒，"
            f"合计{total_time:.3f}秒，测试通过。\n"
            f"结果目录：{case['result_dir']}；训练数据分布目录：{case['distribution_dir']}。"
        )
    run = first.add_run(result_text)
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.size = Pt(10.5)
    if chart is not None:
        picture_paragraph = cell.add_paragraph()
        picture_paragraph.style = "表格正文"
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Keep the chart fully inside the result cell so the right-side legend
        # is not clipped by Word's table boundary after pagination/rendering.
        picture_paragraph.add_run().add_picture(str(chart), width=Cm(12.8))
        caption = cell.add_paragraph()
        caption.style = "表格正文"
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = caption.add_run(
            f"图{int(case['id'].split('-')[1]) - 1}  {case['dataset']}数据集{case['model']}模型各方法第0—99轮平均测试准确率")
        cap_run.font.name = LATIN_FONT
        cap_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
        cap_run.font.size = Pt(10.5)


def build_report() -> None:
    outline = Document(OUTLINE)
    report = Document(TEMPLATE)

    # Environment table comes directly from the final outline.
    replace_table(report.tables[0], outline.tables[0]._element)

    # Remove DomainNet from the dataset description and renumber remaining items.
    paragraphs = report.paragraphs
    domain_index = next(
        i for i, p in enumerate(paragraphs)
        if p.text.strip() == "(2) DomainNet数据集")
    domain_head = paragraphs[domain_index]
    domain_desc = paragraphs[domain_index + 1]
    remove_paragraph(domain_desc)
    remove_paragraph(domain_head)
    for p in report.paragraphs:
        if p.text.strip().startswith("(3) Multi-Domain Sentiment"):
            p.text = p.text.replace("(3)", "(2)", 1)
        elif p.text.strip().startswith("(4) Multi-Domain Digits"):
            p.text = p.text.replace("(4)", "(3)", 1)

    # Test-item table: exact T-1..T-6 rows from the final outline.
    test_item_xml = copy.deepcopy(outline.tables[1]._element)
    test_item_rows = test_item_xml.findall(qn("w:tr"))
    for row in test_item_rows[7:]:
        test_item_xml.remove(row)
    replace_table(report.tables[1], test_item_xml)

    # Summary table: retain style, remove old T-17/T-18 rows, and update T-1..T-6.
    summary = report.tables[2]
    while len(summary.rows) > 7:
        summary._tbl.remove(summary.rows[-1]._tr)
    test_project = {
        "T-1": "深度学习模型跨域协同平台性能测试",
        "T-2": "深度学习模型跨域协同平台对智能模型的支持与跨域协同准确率测试",
        "T-3": "深度学习模型跨域协同平台对智能模型的支持与跨域协同准确率测试",
        "T-4": "深度学习模型跨域协同平台对智能模型的支持与跨域协同准确率测试",
        "T-5": "深度学习模型跨域协同平台对智能模型的支持与跨域协同准确率测试",
        "T-6": "深度学习模型跨域协同平台对智能模型的支持与跨域协同准确率测试",
    }
    for row_index, case in enumerate(CASES, start=1):
        row = summary.rows[row_index]
        if case["id"] == "T-1":
            result = case["summary"]
        else:
            acc = case["accuracies"]
            gains = case["gains"]
            result = (
                f"FedAvg {acc['FedAvg'] * 100:.2f}%、FedProx {acc['FedProx'] * 100:.2f}%、"
                f"平台{acc['平台'] * 100:.2f}%，平台分别提升{gains[0]:.2f}%和{gains[1]:.2f}%。")
        values = [test_project[case["id"]], case["title"], result, "通过", "结果及证据见对应测试记录。"]
        for c, value in enumerate(values):
            set_cell_text(row.cells[c], value)

    # Use current report result/personnel rows as formatting templates.
    old_case_tables = report.tables[3:9]
    result_row_template = copy.deepcopy(find_result_row(old_case_tables[0]))
    personnel_row_template = copy.deepcopy(find_person_row(old_case_tables[0]))

    # Remove the old case records (including obsolete T-17/T-18).
    record_heading = next(p for p in report.paragraphs if p.text.strip() == "测试记录")
    body = report._element.body
    start_index = list(body).index(record_heading._p) + 1
    sect_pr = body.sectPr
    for element in list(body)[start_index:]:
        if element is not sect_pr:
            body.remove(element)

    # Find final-outline headings and tables in order and copy them verbatim.
    outline_headings = {}
    for p in outline.paragraphs:
        text = p.text.strip()
        for case in CASES:
            if text == case["title"]:
                outline_headings[case["id"]] = p._p
    outline_case_tables = outline.tables[2:8]
    chart_paths = {case["id"]: make_chart(case) for case in CASES[1:]}
    inserted_tables = []
    for case, source_table in zip(CASES, outline_case_tables):
        heading_xml = copy.deepcopy(outline_headings[case["id"]])
        table_xml = copy.deepcopy(source_table._element)
        table_xml.append(copy.deepcopy(result_row_template))
        table_xml.append(copy.deepcopy(personnel_row_template))
        body.insert(len(body) - 1, heading_xml)
        body.insert(len(body) - 1, table_xml)
        inserted_tables.append(Table(table_xml, report))

    # Fill the newly appended result/personnel rows and attach charts.
    for case, table in zip(CASES, inserted_tables):
        result_row = table.rows[-2]
        personnel_row = table.rows[-1]
        set_cell_text(result_row.cells[0], "实际结果")
        set_report_result_cell(
            result_row.cells[1], case, chart_paths.get(case["id"]))
        set_cell_text(personnel_row.cells[0], "用例执行人员")
        set_cell_text(personnel_row.cells[1], "项目测试组")
        set_cell_text(personnel_row.cells[2], "执行日期：2026.08.19")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    report.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
